# Copyright (c) 2026, Frappe Technologies Pvt. Ltd. and Contributors
# License: GNU General Public License v3. See license.txt

"""Salary Bank Statement.

A ready-to-hand-to-the-bank salary disbursement statement built straight from
submitted Salary Slips + the Employee master — NOT from the Payroll Entry's
payment_account (which the stock Bank Remittance report requires and which is
often unset, leaving it empty).

Columns match the firm's bank statement layout:
    S. No. | Name | Designation | Bank | A/C No. | IFSC | Amount(net pay)

Bank name / account come from the slip if it stored them, else the Employee;
IFSC and designation come from the Employee. `download_salary_bank_statement`
emits the exact per-company workbook (company header + total row) as .xlsx.
"""

import frappe
from frappe import _
from frappe.utils import flt, formatdate, getdate


def execute(filters=None):
	filters = frappe._dict(filters or {})
	if not filters.get("company"):
		frappe.throw(_("Please select a Company."))
	return get_columns(filters), get_rows(filters)


def get_columns(filters):
	third = _("Father / Husband Name") if filters.get("use_father_husband_name") else _("Designation")
	return [
		{"label": _("S. No."), "fieldname": "sno", "fieldtype": "Int", "width": 60},
		{"label": _("Name"), "fieldname": "employee_name", "fieldtype": "Data", "width": 200},
		{"label": third, "fieldname": "designation", "fieldtype": "Data", "width": 160},
		{"label": _("Bank"), "fieldname": "bank_name", "fieldtype": "Data", "width": 180},
		{"label": _("A/C No."), "fieldname": "account_no", "fieldtype": "Data", "width": 160},
		{"label": _("IFSC"), "fieldname": "ifsc", "fieldtype": "Data", "width": 120},
		{"label": _("Amount"), "fieldname": "amount", "fieldtype": "Currency", "width": 120},
	]


def get_rows(filters):
	use_fh = bool(filters.get("use_father_husband_name"))
	# Bank = show only this bank; Exclude Bank = leave this bank out (e.g. pay
	# same-bank transfers in one file and everyone else by NEFT in another).
	only_bank = (filters.get("bank") or "").strip()
	exclude_bank = (filters.get("exclude_bank") or "").strip()

	slips = _get_slips(filters)
	emp_map = _employee_map([s.employee for s in slips])

	rows = []
	sno = 0
	for s in slips:
		emp = emp_map.get(s.employee, {})
		row_bank = s.bank_name or emp.get("bank_name") or ""
		if only_bank and row_bank != only_bank:
			continue
		if exclude_bank and row_bank == exclude_bank:
			continue
		sno += 1
		rows.append(
			{
				"sno": sno,
				"employee_name": s.employee_name or emp.get("employee_name"),
				"designation": (emp.get("father_or_husband_name") if use_fh else emp.get("designation")) or "",
				"bank_name": row_bank,
				"account_no": s.bank_account_no or emp.get("bank_ac_no") or "",
				"ifsc": emp.get("ifsc_code") or "",
				"amount": flt(s.net_pay),
			}
		)
	return rows


def _get_slips(filters):
	conds = [["docstatus", "=", 1], ["company", "=", filters.company]]
	if filters.get("from_date"):
		conds.append(["start_date", ">=", filters.from_date])
	if filters.get("to_date"):
		conds.append(["start_date", "<=", filters.to_date])
	return frappe.get_all(
		"Salary Slip",
		filters=conds,
		fields=["employee", "employee_name", "bank_name", "bank_account_no", "net_pay", "start_date"],
		order_by="employee_name asc",
	)


def _employee_map(employees):
	if not employees:
		return {}
	rows = frappe.get_all(
		"Employee",
		filters={"name": ["in", list(set(employees))]},
		fields=[
			"name",
			"employee_name",
			"designation",
			"bank_name",
			"bank_ac_no",
			"ifsc_code",
			"father_or_husband_name",
		],
	)
	return {r.name: r for r in rows}


@frappe.whitelist()
def download_salary_bank_statement(
	company, from_date=None, to_date=None, use_father_husband_name=0, bank=None, exclude_bank=None
):
	"""Emit the exact per-company bank statement workbook (.xlsx): company header,
	period title, the columns, and a TOTAL row — the firm's ready-to-upload format."""
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font

	filters = frappe._dict(
		{
			"company": company,
			"from_date": from_date,
			"to_date": to_date,
			"use_father_husband_name": frappe.utils.cint(use_father_husband_name),
			"bank": bank,
			"exclude_bank": exclude_bank,
		}
	)
	rows = get_rows(filters)
	third_label = "Father / Husband Name" if filters.use_father_husband_name else "Designation"

	wb = Workbook()
	ws = wb.active
	ws.title = (company or "Bank")[:28]

	bold = Font(bold=True)
	center = Alignment(horizontal="center")

	# Header block
	ws.append([company])
	ws["A1"].font = Font(bold=True, size=13)
	addr = _company_address(company)
	if addr:
		ws.append([addr])
	period = _period_label(from_date, to_date)
	ws.append([f"Salary Statement for the month of {period}" if period else "Salary Statement"])
	ws.append([])

	header = ["S. No.", "Name", third_label, "Bank", "A/C No.", "IFSC", "Amount"]
	ws.append(header)
	for c in ws[ws.max_row]:
		c.font = bold

	total = 0.0
	for r in rows:
		ws.append(
			[r["sno"], r["employee_name"], r["designation"], r["bank_name"], r["account_no"], r["ifsc"], r["amount"]]
		)
		total += flt(r["amount"])

	ws.append(["", "", "", "", "", "TOTAL", total])
	last = ws[ws.max_row]
	last[5].font = bold
	last[6].font = bold

	# Keep account numbers as text (avoid scientific notation / dropped leading zeros)
	for row in ws.iter_rows(min_row=6, min_col=5, max_col=5):
		for cell in row:
			cell.number_format = "@"
			if cell.value is not None:
				cell.value = str(cell.value)

	widths = [8, 30, 24, 26, 22, 16, 14]
	for idx, w in enumerate(widths, start=1):
		ws.column_dimensions[chr(64 + idx)].width = w

	import io

	buf = io.BytesIO()
	wb.save(buf)
	fname = f"Salary_Bank_Statement_{(company or '').replace(' ', '_')}_{period or ''}.xlsx".replace("__", "_")
	frappe.response["filename"] = fname
	frappe.response["filecontent"] = buf.getvalue()
	frappe.response["type"] = "binary"


def _period_label(from_date, to_date):
	d = from_date or to_date
	if not d:
		return ""
	return formatdate(getdate(d), "MMMM yyyy")


def _company_address(company):
	"""Best-effort single-line address for the company header — blank if none."""
	try:
		link = frappe.get_all(
			"Dynamic Link",
			filters={"link_doctype": "Company", "link_name": company, "parenttype": "Address"},
			fields=["parent"],
			limit=1,
		)
		if not link:
			return ""
		a = frappe.db.get_value(
			"Address", link[0].parent, ["address_line1", "city", "state"], as_dict=True
		)
		if not a:
			return ""
		return ", ".join(p for p in [a.address_line1, a.city, a.state] if p)
	except Exception:
		return ""
